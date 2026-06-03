#!/usr/bin/env python3
"""Convert morphometric landmark data between TPS / NTSys / CSV / JSON / Excel."""
from __future__ import annotations

import argparse
import json
import os
import re
import sys
from pathlib import Path

_HERE = os.path.dirname(os.path.abspath(__file__))
_SHARED = os.path.normpath(os.path.join(_HERE, "..", "..", "scripts"))
if _SHARED not in sys.path:
    sys.path.insert(0, _SHARED)

from paleo_common import print_json  # noqa: E402

# ========== Parsers ==========

def _parse_tps(text):
    """Return list of specimen dicts."""
    specimens = []
    cur = None
    lines = text.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].strip()
        if not line:
            i += 1
            continue
        if line.upper().startswith("LM="):
            if cur and cur["landmarks"]:
                specimens.append(cur)
            n = int(line.split("=", 1)[1])
            cur = {"specimen_id": f"specimen_{len(specimens)+1:03d}",
                   "scale": 1.0,
                   "image_path": None,
                   "num_landmarks": n,
                   "landmarks": []}
            i += 1
            for k in range(n):
                if i >= len(lines):
                    break
                parts = re.split(r"[\s,]+", lines[i].strip())
                if len(parts) >= 2:
                    try:
                        x = float(parts[0]); y = float(parts[1])
                    except ValueError:
                        i += 1; continue
                    cur["landmarks"].append({
                        "id": k + 1, "x": x, "y": y,
                        "type": "semilandmark", "name": f"LM{k+1:02d}",
                    })
                i += 1
            continue
        if cur is not None:
            if line.upper().startswith("ID="):
                cur["specimen_id"] = line.split("=", 1)[1].strip()
            elif line.upper().startswith("IMAGE="):
                cur["image_path"] = line.split("=", 1)[1].strip()
            elif line.upper().startswith("SCALE="):
                try:
                    cur["scale"] = float(line.split("=", 1)[1].strip())
                except ValueError:
                    pass
        i += 1
    if cur and cur["landmarks"]:
        specimens.append(cur)
    return specimens


def _parse_csv(text):
    """Parse long-format CSV: specimen_id,landmark_id,x,y[,type,name]."""
    lines = [l for l in text.splitlines() if l.strip()]
    if not lines:
        return []
    header = [h.strip().lower() for h in lines[0].split(",")]
    sid_idx = header.index("specimen_id") if "specimen_id" in header else 0
    lmid_idx = header.index("landmark_id") if "landmark_id" in header else 1
    x_idx = header.index("x") if "x" in header else 2
    y_idx = header.index("y") if "y" in header else 3
    type_idx = header.index("type") if "type" in header else None
    name_idx = header.index("name") if "name" in header else None

    grouped = {}
    for line in lines[1:]:
        parts = [p.strip() for p in line.split(",")]
        if len(parts) < 4:
            continue
        sid = parts[sid_idx]
        try:
            lmid = int(parts[lmid_idx])
            x = float(parts[x_idx])
            y = float(parts[y_idx])
        except ValueError:
            continue
        ent = grouped.setdefault(sid, [])
        ent.append({
            "id": lmid, "x": x, "y": y,
            "type": parts[type_idx] if type_idx is not None and type_idx < len(parts) else "semilandmark",
            "name": parts[name_idx] if name_idx is not None and name_idx < len(parts) else f"LM{lmid:02d}",
        })
    out = []
    for sid, lms in grouped.items():
        lms.sort(key=lambda x: x["id"])
        out.append({"specimen_id": sid, "scale": 1.0, "image_path": None,
                    "num_landmarks": len(lms), "landmarks": lms})
    return out


def _parse_json(text):
    data = json.loads(text)
    if isinstance(data, dict) and "landmarks" in data:
        return [data]
    if isinstance(data, list):
        return [d for d in data if isinstance(d, dict) and "landmarks" in d]
    return []


def _parse_nts(text):
    """NTSys matrix: header `1 N P 0` then N rows of P columns."""
    lines = [l for l in text.splitlines() if l.strip() and not l.startswith('"')]
    if len(lines) < 2:
        return []
    header = re.split(r"\s+", lines[0].strip())
    try:
        n_specimens = int(header[1])
        n_vars = int(header[2])
    except (ValueError, IndexError):
        return []
    n_landmarks = n_vars // 2
    rows = lines[1:1 + n_specimens]
    specimens = []
    for i, row in enumerate(rows):
        vals = re.split(r"\s+", row.strip())
        lms = []
        for k in range(n_landmarks):
            try:
                x = float(vals[2 * k])
                y = float(vals[2 * k + 1])
            except (ValueError, IndexError):
                break
            lms.append({"id": k + 1, "x": x, "y": y,
                        "type": "semilandmark", "name": f"LM{k+1:02d}"})
        specimens.append({"specimen_id": f"specimen_{i+1:03d}", "scale": 1.0,
                          "image_path": None, "num_landmarks": len(lms),
                          "landmarks": lms})
    return specimens


def _parse_xlsx(path):
    try:
        from openpyxl import load_workbook
    except ImportError:
        return None, "openpyxl not installed (pip install openpyxl)"
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return [], None
    header = [str(h or "").strip().lower() for h in rows[0]]
    body = ["," .join("" if v is None else str(v) for v in r) for r in rows[1:]]
    csv_text = ",".join(header) + "\n" + "\n".join(body)
    return _parse_csv(csv_text), None


# ========== Writers ==========

def _write_tps(specimens, out_path):
    lines = []
    for s in specimens:
        lines.append(f"LM={len(s['landmarks'])}")
        for lm in s["landmarks"]:
            lines.append(f"{lm['x']:.4f}  {lm['y']:.4f}")
        if s.get("image_path"):
            lines.append(f"IMAGE={Path(s['image_path']).name}")
        lines.append(f"ID={s['specimen_id']}")
        lines.append(f"SCALE={s.get('scale', 1.0):.4f}")
        lines.append("")
    Path(out_path).parent.mkdir(parents=True, exist_ok=True)
    Path(out_path).write_text("\n".join(lines), encoding="utf-8")


def _write_csv(specimens, out_path):
    rows = ["specimen_id,landmark_id,x,y,type,name"]
    for s in specimens:
        sid = s["specimen_id"]
        for lm in s["landmarks"]:
            rows.append(f"{sid},{lm['id']},{lm['x']:.4f},{lm['y']:.4f},"
                        f"{lm['type']},{lm['name']}")
    Path(out_path).parent.mkdir(parents=True, exist_ok=True)
    Path(out_path).write_text("\n".join(rows) + "\n", encoding="utf-8")


def _write_json(specimens, out_path):
    Path(out_path).parent.mkdir(parents=True, exist_ok=True)
    payload = specimens[0] if len(specimens) == 1 else specimens
    Path(out_path).write_text(json.dumps(payload, ensure_ascii=False, indent=2),
                              encoding="utf-8")


def _write_nts(specimens, out_path):
    if not specimens:
        Path(out_path).write_text("", encoding="utf-8")
        return
    n = len(specimens)
    p_var = max(len(s["landmarks"]) for s in specimens) * 2
    lines = [f"1 {n} {p_var} 0"]
    for s in specimens:
        flat = []
        for lm in s["landmarks"]:
            flat.extend([f"{lm['x']:.4f}", f"{lm['y']:.4f}"])
        lines.append(" ".join(flat))
    Path(out_path).parent.mkdir(parents=True, exist_ok=True)
    Path(out_path).write_text("\n".join(lines) + "\n", encoding="utf-8")


def _write_xlsx(specimens, out_path):
    try:
        from openpyxl import Workbook
    except ImportError:
        return False, "openpyxl not installed (pip install openpyxl)"
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws.append(["specimen_id", "landmark_id", "x", "y", "type", "name", "scale"])
    for s in specimens:
        sid = s["specimen_id"]
        scale = s.get("scale", 1.0)
        for lm in s["landmarks"]:
            ws.append([sid, lm["id"], lm["x"], lm["y"], lm["type"], lm["name"], scale])
    Path(out_path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    return True, None


PARSERS = {".tps": _parse_tps, ".csv": _parse_csv, ".json": _parse_json,
           ".nts": _parse_nts, ".txt": _parse_tps}
WRITERS = {"tps": _write_tps, "csv": _write_csv, "json": _write_json,
           "nts": _write_nts}


def _load_one(path):
    p = Path(path)
    ext = p.suffix.lower()
    if ext in (".xlsx", ".xls"):
        return _parse_xlsx(p)
    fn = PARSERS.get(ext)
    if not fn:
        return None, f"Unsupported input extension: {ext}"
    try:
        text = p.read_text(encoding="utf-8")
    except OSError as e:
        return None, f"Failed to read {p}: {e}"
    return fn(text), None


def _validate(specimens):
    if not specimens:
        return {"all_same_count": True, "missing_landmarks": 0,
                "warnings": ["No specimens"]}
    counts = [len(s["landmarks"]) for s in specimens]
    same = len(set(counts)) == 1
    warnings = []
    if not same:
        warnings.append(f"Inconsistent landmark counts: {sorted(set(counts))}")
    missing = sum(1 for s in specimens for lm in s["landmarks"]
                  if lm["x"] is None or lm["y"] is None)
    return {
        "all_same_count": same,
        "landmark_counts": sorted(set(counts)),
        "missing_landmarks": missing,
        "specimens_count": len(specimens),
        "warnings": warnings,
    }


def main(argv=None):
    p = argparse.ArgumentParser(description="Convert morphometric landmark data between formats")
    p.add_argument("--input", help="Input file (or use --merge for multiple)")
    p.add_argument("--merge", nargs="+", help="Merge multiple input files")
    p.add_argument("--to", help="Target format: tps, csv, json, nts, xlsx")
    p.add_argument("--out", help="Output path")
    p.add_argument("--validate", action="store_true",
                   help="Validate input only, no conversion")
    args = p.parse_args(argv)

    inputs = [args.input] if args.input else (args.merge or [])
    if not inputs:
        print_json({"error": "Provide --input or --merge"})
        return

    all_specimens = []
    in_formats = []
    for inp in inputs:
        spec, err = _load_one(inp)
        if err:
            print_json({"error": err, "input": inp})
            return
        all_specimens.extend(spec or [])
        in_formats.append(Path(inp).suffix.lower().lstrip("."))

    validation = _validate(all_specimens)

    if args.validate:
        print_json({
            "inputs": inputs,
            "input_formats": in_formats,
            "validation": validation,
            "specimens_preview": [
                {"specimen_id": s["specimen_id"], "num_landmarks": len(s["landmarks"])}
                for s in all_specimens[:10]
            ],
        })
        return

    if not args.to or not args.out:
        print_json({"error": "Provide --to and --out (or use --validate only)",
                    "validation": validation})
        return

    fmt = args.to.lower()
    if fmt in ("xlsx", "excel"):
        ok, err = _write_xlsx(all_specimens, args.out)
        if not ok:
            print_json({"error": err})
            return
    elif fmt in WRITERS:
        WRITERS[fmt](all_specimens, args.out)
    else:
        print_json({"error": f"Unsupported output format: {fmt}",
                    "supported": list(WRITERS) + ["xlsx"]})
        return

    print_json({
        "inputs": inputs,
        "input_formats": in_formats,
        "output_format": fmt,
        "output_path": args.out,
        "specimens_count": len(all_specimens),
        "landmarks_per_specimen": validation.get("landmark_counts"),
        "validation": validation,
    })


if __name__ == "__main__":
    main()
